"""Read InChI species data and export SMILES and explicit-H graph workbooks.

Importing this module only defines the class. Run run_species_readout.py to
create the two workbooks. CSV identities are preserved, never silently corrected.
"""
import csv
import io
import json
import math
from collections import Counter
from pathlib import Path

from PIL import Image as PILImage, ImageDraw, ImageFont

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from rdkit import Chem, rdBase
from rdkit.Chem import rdDepictor


class SpeciesReadOut:
    """One CSV reader shared by both exports.

    The CSV is the only input. Output paths are supplied to the export methods
    and remain relative if given that way.
    """

    def __init__(self, csv_path):
        self.csv_path = Path(csv_path)
        self.records = []
        with self.csv_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != ["Species", "Chemical Name", "InChI"]:
                raise ValueError("CSV columns must be: Species, Chemical Name, InChI")
            seen = set()
            for line, row in enumerate(reader, 2):
                if None in row or any(not isinstance(v, str) or not v.strip() for v in row.values()):
                    raise ValueError(f"CSV row {line}: missing values or extra fields")
                key = row["Species"]
                if key in seen:
                    raise ValueError(f"CSV row {line}: duplicate species {key}")
                seen.add(key)
                try:
                    self.records.append(self.convert(row))
                except (ValueError, RuntimeError) as exc:
                    raise ValueError(f"CSV row {line}, species {key}: {exc}") from exc
        if not self.records:
            raise ValueError("CSV contains no species")
        print(f"Read {len(self.records)} species from {self.csv_path}; RDKit {rdBase.rdkitVersion}")

    @staticmethod
    def compact(value):
        """Compact arrays are easier to copy into a Python/JSON graph definition."""
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

    @staticmethod
    def split_state(inchi):
        """Remove only recognized table annotations, retaining their meaning.

        These prefixes are not InChI syntax. RDKit's parsed radical-electron
        count and SMILES must not be treated as a recovered electronic state.
        """
        for prefix, state in (("EXCITED_", "Excited; see source chemical name"),
                              ("SINGLET_", "Singlet")):
            if inchi.startswith(prefix):
                return inchi[len(prefix):], state
        return inchi, "Not encoded; see source chemical name"

    def convert(self, source):
        key = source["Species"]
        normalized, state = self.split_state(source["InChI"].strip())
        if not normalized.startswith("InChI="):
            raise ValueError("Expected InChI=... or a recognized EXCITED_/SINGLET_ prefix")
        # Retain standalone H atoms; they are valid species in combustion data.
        molecule = Chem.MolFromInchi(normalized, sanitize=True, removeHs=False)
        if molecule is None or molecule.GetNumAtoms() == 0:
            raise ValueError("RDKit could not parse the InChI")
        smiles_molecule = (Chem.RemoveHs(molecule) if any(a.GetAtomicNum() != 1
                           for a in molecule.GetAtoms()) else molecule)
        smiles = Chem.MolToSmiles(smiles_molecule, canonical=True, isomericSmiles=True)

        # Add every hydrogen explicitly, then put all base atoms first.
        explicit = Chem.AddHs(molecule)
        Chem.Kekulize(explicit, clearAromaticFlags=True)
        heavy = [a.GetIdx() for a in explicit.GetAtoms() if a.GetAtomicNum() != 1]
        hydrogens = [a.GetIdx() for a in explicit.GetAtoms() if a.GetAtomicNum() == 1]
        explicit = Chem.RenumberAtoms(explicit, heavy + hydrogens)
        symbols = [a.GetSymbol() for a in explicit.GetAtoms()]
        # H and H2 have no heavy atoms, so their H atoms are base atoms.
        base_count = len(heavy) if heavy else len(symbols)
        bonds = sorted([min(b.GetBeginAtomIdx(), b.GetEndAtomIdx()),
                        max(b.GetBeginAtomIdx(), b.GetEndAtomIdx()),
                        int(b.GetBondTypeAsDouble())] for b in explicit.GetBonds())
        if any(b.GetBondTypeAsDouble() not in (1, 2, 3) for b in explicit.GetBonds()):
            raise ValueError("Graph export requires single, double, or triple bonds")
        h_counts = ([sum(n.GetAtomicNum() == 1 for n in explicit.GetAtomWithIdx(i).GetNeighbors())
                     for i in range(base_count)] if heavy else [])
        # A disconnected H fragment cannot be represented by attached-H counts.
        if heavy and sum(h_counts) != len(hydrogens):
            raise ValueError("Some hydrogen atoms are not attached to base atoms")
        counts = Counter(symbols)
        family = ("Inert / bath gas" if key in ("HE", "AR", "N2") else
                  f"C{counts['C']} chemistry" if counts["C"] else "H / O chemistry")
        notes = []
        if state.startswith("Excited") or state == "Singlet":
            notes.append(f"{state}. SMILES does not encode this state.")
        if not notes:
            notes.append("CSV InChI connectivity; electronic state is not inferred.")
        return {
            "species": key, "chemical_name": source["Chemical Name"],
            "inchi": source["InChI"], "normalized_inchi": normalized, "smiles": smiles,
            "state": state, "notes": " ".join(notes), "family": family,
            "base_atoms": symbols[:base_count],
            "base_bonds": [b for b in bonds if b[1] < base_count], "h_counts": h_counts,
            "symbols": symbols, "bonds": bonds,
            "charges": [a.GetFormalCharge() for a in explicit.GetAtoms()],
            "rdkit_radical_electrons": [a.GetNumRadicalElectrons() for a in explicit.GetAtoms()],
            "molecule": explicit,
        }

    @staticmethod
    def write_text_row(sheet, row, values):
        # Store identifiers literally, including strings that resemble formulas.
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = "s"

    @staticmethod
    def style_table(sheet, header_row, widths, row_height, table_name):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = f"B{header_row + 1}"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for cells in sheet.iter_rows(min_row=header_row):
            row = cells[0].row
            sheet.row_dimensions[row].height = 32 if row == header_row else row_height
            for cell in cells:
                cell.font = Font(name="Arial", size=11, color="17364B")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row == header_row:
                    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173F56")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif (row - header_row) % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F7FA")
        table = Table(displayName=table_name,
                      ref=f"A{header_row}:{get_column_letter(len(widths))}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)

    def save(self, workbook, output_path):
        path = Path(output_path)
        if path.resolve() == self.csv_path.resolve():
            raise ValueError("Output path would overwrite an input file")
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Output filename must end in .xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()
        print(f"Saved {path} ({len(self.records)} species)")
        return path

    def write_smiles_workbook(self, output_path):
        """First export: exactly the source's three columns plus SMILES."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Species with SMILES"
        self.write_text_row(sheet, 1, ["Species", "Chemical Name", "InChI", "SMILES"])
        for row, record in enumerate(self.records, 2):
            self.write_text_row(sheet, row, [record[k] for k in
                                           ("species", "chemical_name", "inchi", "smiles")])
        self.style_table(sheet, 1, [17, 53, 76, 34], 44, "SpeciesSmiles")
        return self.save(workbook, output_path)

    @staticmethod
    def structure_png(record):
        """Draw bonds behind labeled atom circles using RDKit's 2D coordinates."""
        molecule = Chem.Mol(record["molecule"])
        rdDepictor.Compute2DCoords(molecule)
        coordinates = molecule.GetConformer()
        xy = [(coordinates.GetAtomPosition(i).x, coordinates.GetAtomPosition(i).y)
              for i in range(molecule.GetNumAtoms())]
        xmin, xmax = min(x for x, _ in xy), max(x for x, _ in xy)
        ymin, ymax = min(y for _, y in xy), max(y for _, y in xy)
        # Equal scaling preserves the layout; padding includes the atom circles.
        width, height, padding = 840, 400, 48
        scale = min(120, (width - 2 * padding) / max(xmax - xmin, 1e-6),
                    (height - 2 * padding) / max(ymax - ymin, 1e-6))
        points = [(width / 2 + (x - (xmin + xmax) / 2) * scale,
                   height / 2 - (y - (ymin + ymax) / 2) * scale) for x, y in xy]
        nearest = min((math.dist(a, b) for i, a in enumerate(points)
                       for b in points[i+1:]), default=100)
        radius = min(32, nearest * 0.43)
        canvas = PILImage.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(canvas)
        # Parallel strokes make single/double/triple bonds explicit. Opaque
        # circles drawn afterwards cover the bond ends beneath each atom.
        for bond in molecule.GetBonds():
            a, b = points[bond.GetBeginAtomIdx()], points[bond.GetEndAtomIdx()]
            length = math.dist(a, b)
            nx, ny = -(b[1] - a[1]) / length, (b[0] - a[0]) / length
            order = int(bond.GetBondTypeAsDouble())
            for stroke in range(order):
                shift = (stroke - (order - 1) / 2) * 8
                draw.line((a[0] + nx * shift, a[1] + ny * shift,
                           b[0] + nx * shift, b[1] + ny * shift), fill="#66788A", width=3)
        colors = {"C": "#DEF0FC", "H": "#FFFFFF", "O": "#FFE2DB",
                  "N": "#DBE5FF", "He": "#E9DFF9", "Ar": "#E9DFF9"}
        for atom in molecule.GetAtoms():
            charge = atom.GetFormalCharge()
            charge_text = ("+" if charge == 1 else "-" if charge == -1 else
                           f"{charge:+d}" if charge else "")
            label = f"{atom.GetSymbol()}{atom.GetIdx()}{charge_text}"
            size = 26
            font = ImageFont.load_default(size=size)
            while draw.textlength(label, font=font) > radius * 1.65 and size > 10:
                size -= 1
                font = ImageFont.load_default(size=size)
            x, y = points[atom.GetIdx()]
            draw.ellipse((x-radius, y-radius, x+radius, y+radius),
                         fill=colors.get(atom.GetSymbol(), "#EDF1F4"), outline="#66788A", width=2)
            draw.text((x, y), label, anchor="mm", font=font, fill="#18364B")
        output = io.BytesIO()
        canvas.save(output, format="PNG")
        return output.getvalue()

    def reference_heading(self, sheet, title, lines):
        for row, text in enumerate([title, *lines], 2):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            sheet.cell(row, 1, text)
            sheet.cell(row, 1).font = Font(name="Arial", size=18 if row == 2 else 11,
                                          bold=row == 2, color="173F56")
            sheet.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[row].height = 32 if row == 2 else 28
        sheet.cell(2, 7, f"Source CSV: {self.csv_path}")
        sheet.cell(3, 7, "Structure source: supplied CSV only")
        sheet.cell(4, 7, f"RDKit version: {rdBase.rdkitVersion}")
        for row in range(2, 5):
            sheet.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
            sheet.cell(row, 7).font = Font(name="Arial", size=11, color="17364B")
            sheet.cell(row, 7).alignment = Alignment(vertical="center", wrap_text=True)

    def write_species_reference(self, output_path):
        """Second export: diagrams plus machine-readable graph definitions.

        As in the existing reference, headers are on row 7 and species start
        on row 8. Each diagram is inset inside its own cell, with equal aspect
        ratio and explicit row/column sizes, so adjacent structures cannot overlap.
        """
        workbook = Workbook()
        reference = workbook.active
        reference.title = "Species reference"
        definitions = workbook.create_sheet("Graph definitions")
        lines = [f"{len(self.records)} species derived from CSV InChI values; source identities retained.",
                 "Explicit H atoms; labels = element + zero-based node index; parallel lines = bond order.",
                 "SMILES and RDKit radical counts do not encode the source electronic state or spin multiplicity.",
                 "Base atoms precede appended H. Attached-H counts follow base order; H and H2 use []."]
        for sheet in (reference, definitions):
            self.reference_heading(sheet, "InChI species reference", lines)
        self.write_text_row(reference, 7, ["Species key", "Family", "Structure (all atoms)",
            "Base atom indices", "Base bonds (u, v, order)", "Attached H per base atom",
            "Chemical name", "State / representation notes"])
        headers = ["Species key", "Family", "Base atoms JSON", "Base bonds JSON", "H counts JSON",
                   "All node symbols JSON", "All bonds JSON", "Base formal charges JSON",
                   "Chemical name", "State / representation notes",
                   "Source InChI", "Normalized InChI", "SMILES", "Source state annotation",
                   "All formal charges JSON", "RDKit radical electrons JSON"]
        self.write_text_row(definitions, 7, headers)
        for row, record in enumerate(self.records, 8):
            base_count = len(record["base_atoms"])
            self.write_text_row(reference, row, [record["species"], record["family"], "",
                "\n".join(f"{i}: {s}" for i, s in enumerate(record["base_atoms"])),
                self.compact(record["base_bonds"]), self.compact(record["h_counts"]),
                record["chemical_name"], record["notes"]])
            self.write_text_row(definitions, row, [record["species"], record["family"],
                self.compact(record["base_atoms"]), self.compact(record["base_bonds"]),
                self.compact(record["h_counts"]), self.compact(record["symbols"]),
                self.compact(record["bonds"]), self.compact(record["charges"][:base_count]),
                record["chemical_name"], record["notes"],
                record["inchi"], record["normalized_inchi"], record["smiles"], record["state"],
                self.compact(record["charges"]), self.compact(record["rdkit_radical_electrons"])])
            # Draw at 2x display resolution. Display is 420 x 200 pixels inside
            # a ~453 x 240 pixel cell (column width 64, row height 180 points).
            picture = Image(io.BytesIO(self.structure_png(record)))
            picture.anchor = OneCellAnchor(
                _from=AnchorMarker(col=2, row=row-1, colOff=pixels_to_EMU(12), rowOff=pixels_to_EMU(20)),
                ext=XDRPositiveSize2D(pixels_to_EMU(420), pixels_to_EMU(200)))
            reference.add_image(picture)
        self.style_table(reference, 7, [17, 20, 64, 24, 30, 24, 38, 57], 180, "SpeciesReference")
        self.style_table(definitions, 7, [17, 20, 27, 36, 24, 44, 58, 28, 38, 57,
                                          70, 70, 35, 43, 30, 34], 140, "GraphDefinitions")
        return self.save(workbook, output_path)
