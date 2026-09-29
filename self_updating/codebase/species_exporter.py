"""Convert species keys and SMILES/InChI strings to a reference workbook.

Example (run from self_updating):
    species = [{"species_key": "CH4", "smiles": "C"},
               {"species_key": "H2O", "inchi": "InChI=1S/H2O/h1H2"}]
    exporter = SpeciesExporter(species)
    exporter.write_xlsx("input/my_species.xlsx")

Importing this module creates no files. RDKit interprets the supplied structure;
the species key is a label, not a formula used to invent or change connectivity.
"""
import io
import json
import math
from pathlib import Path

from PIL import Image as PILImage, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from rdkit import Chem, rdBase
from rdkit.Chem import rdDepictor, rdMolDescriptors


class SpeciesExporter:
    """Each input row needs species_key and exactly one of smiles or inchi.

    An optional chemical_name is copied without a lookup. EXCITED_/SINGLET_
    prefixes in source InChI values are retained as annotations, not inferred
    from the RDKit molecule. All entries are
    converted before writing, so a bad input cannot produce a partial table.
    Duplicate keys and disconnected structures are rejected. Other elements
    can be exported, but the current training network still supports only CHO.
    """

    def __init__(self, species):
        self.records = []
        seen = set()
        for row in species:
            if not isinstance(row, dict):
                raise ValueError("Each species must be a dictionary")
            key = row.get("species_key")
            if not isinstance(key, str) or not key.strip():
                raise ValueError("Each species needs a nonempty species_key")
            key = key.strip()
            if key in seen:
                raise ValueError(f"Duplicate species_key: {key}")
            seen.add(key)
            self.records.append(self.convert({**row, "species_key": key}))
        if not self.records:
            raise ValueError("Provide at least one species")

    @staticmethod
    def convert(row):
        """Expand H explicitly, place base atoms first, and keep bond orders."""
        key = row["species_key"]
        sources = {}
        for field in ("smiles", "inchi"):
            value = row.get(field, "")
            if not isinstance(value, str):
                raise ValueError(f"{key}: {field} must be a string")
            if value.strip():
                sources[field] = value.strip()
        if len(sources) != 1:
            raise ValueError(f"{key}: supply exactly one of smiles or inchi")
        kind, structure = next(iter(sources.items()))
        # FFCM2 uses these labels before InChI strings. They are source metadata,
        # not standard InChI syntax, and cannot be recovered from connectivity.
        state_annotation = ""
        if kind == "inchi":
            for prefix, state in (("EXCITED_", "Excited"), ("SINGLET_", "Singlet")):
                if structure.startswith(prefix):
                    state_annotation = state
                    structure = structure[len(prefix):]
                    break
        # Report invalid inputs with the species key instead of RDKit's timed log.
        with rdBase.BlockLogs():
            if kind == "smiles":
                options = Chem.SmilesParserParams()
                options.removeHs = False
                options.parseName = False
                mol = Chem.MolFromSmiles(structure, options)
            else:
                mol = Chem.MolFromInchi(structure, sanitize=True, removeHs=False)
        if mol is None or mol.GetNumAtoms() == 0:
            raise ValueError(f"{key}: invalid {kind} structure {structure!r}")
        if len(Chem.GetMolFrags(mol)) != 1:
            raise ValueError(f"{key}: provide one connected molecule, not disconnected fragments")
        if any(a.GetAtomicNum() == 0 for a in mol.GetAtoms()):
            raise ValueError(f"{key}: wildcard atoms do not define a species")

        with rdBase.BlockLogs():
            smiles = Chem.MolToSmiles(Chem.RemoveHs(mol), isomericSmiles=True)
        formula = rdMolDescriptors.CalcMolFormula(mol)
        explicit = Chem.AddHs(mol)
        # The graph reader expects integer orders, including for aromatic input.
        Chem.Kekulize(explicit, clearAromaticFlags=True)
        heavy = [a.GetIdx() for a in explicit.GetAtoms() if a.GetAtomicNum() != 1]
        hydrogen = [a.GetIdx() for a in explicit.GetAtoms() if a.GetAtomicNum() == 1]
        explicit = Chem.RenumberAtoms(explicit, heavy + hydrogen)
        base_count = len(heavy) or len(hydrogen)  # H and H2 are themselves base atoms.
        symbols = [a.GetSymbol() for a in explicit.GetAtoms()]
        charges = [a.GetFormalCharge() for a in explicit.GetAtoms()]
        bonds = []
        for bond in explicit.GetBonds():
            order = bond.GetBondTypeAsDouble()
            if order not in (1.0, 2.0, 3.0):
                raise ValueError(f"{key}: unsupported bond order {order}")
            u, v = sorted((bond.GetBeginAtomIdx(), bond.GetEndAtomIdx()))
            bonds.append([u, v, int(order)])
        bonds.sort()
        h_counts = ([sum(a.GetAtomicNum() == 1 for a in explicit.GetAtomWithIdx(i).GetNeighbors())
                     for i in range(base_count)] if heavy else [])
        notes = []
        if state_annotation:
            notes.append(f"Source state: {state_annotation}. The parsed InChI supplies connectivity; "
                         "SMILES and RDKit radical counts do not establish this electronic state.")
        if any(a.GetIsotope() for a in explicit.GetAtoms()):
            notes.append("Isotopes are retained in SMILES and the plot; the graph reader ignores isotope labels.")
        if (any(a.GetChiralTag() != Chem.ChiralType.CHI_UNSPECIFIED for a in mol.GetAtoms())
                or any(b.GetStereo() != Chem.BondStereo.STEREONONE for b in mol.GetBonds())):
            notes.append("Stereochemistry is retained in SMILES and the plot; the graph reader uses connectivity only.")
        return {
            "species_key": key, "chemical_name": row.get("chemical_name", ""),
            "source_smiles": sources.get("smiles", ""), "source_inchi": sources.get("inchi", ""),
            "normalized_inchi": structure if kind == "inchi" else "",
            "state_annotation": state_annotation,
            "smiles": smiles, "formula": formula, "molecule": explicit,
            "base_atoms": symbols[:base_count], "base_bonds": [b for b in bonds if b[1] < base_count],
            "h_counts": h_counts, "symbols": symbols, "bonds": bonds,
            "charges": charges, "radicals": [a.GetNumRadicalElectrons() for a in explicit.GetAtoms()],
            "notes": " ".join(notes),
        }

    @staticmethod
    def structure_png(record):
        """Draw the explicit graph in the same style as the CHO reference sheet."""
        molecule = Chem.Mol(record["molecule"])
        rdDepictor.Compute2DCoords(molecule)
        conformer = molecule.GetConformer()
        coordinates = [(conformer.GetAtomPosition(index).x, conformer.GetAtomPosition(index).y)
                       for index in range(molecule.GetNumAtoms())]
        xmin, xmax = min(x for x, _ in coordinates), max(x for x, _ in coordinates)
        ymin, ymax = min(y for _, y in coordinates), max(y for _, y in coordinates)
        width, height, padding = 840, 400, 48
        scale = min(120, (width - 2 * padding) / max(xmax - xmin, 1e-6),
                    (height - 2 * padding) / max(ymax - ymin, 1e-6))
        points = [(width / 2 + (x - (xmin + xmax) / 2) * scale,
                   height / 2 - (y - (ymin + ymax) / 2) * scale)
                  for x, y in coordinates]
        nearest = min((math.dist(a, b) for index, a in enumerate(points) for b in points[index + 1:]),
                      default=100)
        radius = min(32, nearest * 0.43)
        canvas = PILImage.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(canvas)
        # Bonds go first. The opaque atom circles cover their end points.
        for bond in molecule.GetBonds():
            a, b = points[bond.GetBeginAtomIdx()], points[bond.GetEndAtomIdx()]
            length = math.dist(a, b)
            normal_x, normal_y = -(b[1] - a[1]) / length, (b[0] - a[0]) / length
            for stroke in range(int(bond.GetBondTypeAsDouble())):
                shift = (stroke - (bond.GetBondTypeAsDouble() - 1) / 2) * 8
                draw.line((a[0] + normal_x * shift, a[1] + normal_y * shift,
                           b[0] + normal_x * shift, b[1] + normal_y * shift),
                          fill="#66788A", width=3)
        colors = {"C": "#DEF0FC", "H": "#FFFFFF", "O": "#FFE2DB",
                  "N": "#DBE5FF", "He": "#E9DFF9", "Ar": "#E9DFF9"}
        for atom in molecule.GetAtoms():
            charge = atom.GetFormalCharge()
            charge_text = ("+" if charge == 1 else "-" if charge == -1 else
                           f"{charge:+d}" if charge else "")
            label = f"{atom.GetSymbol()}{atom.GetIdx()}{charge_text}"
            size = 26
            font = ImageFont.load_default(size=size)
            while draw.textlength(label, font=font) > radius * 1.65 and size > 10:
                size -= 1
                font = ImageFont.load_default(size=size)
            x, y = points[atom.GetIdx()]
            draw.ellipse((x - radius, y - radius, x + radius, y + radius),
                         fill=colors.get(atom.GetSymbol(), "#EDF1F4"),
                         outline="#66788A", width=2)
            draw.text((x, y), label, anchor="mm", font=font, fill="#18364B")
        output = io.BytesIO()
        canvas.save(output, format="PNG")
        return output.getvalue()

    @staticmethod
    def write_row(sheet, row, values):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Species keys/source strings are literal text.

    @staticmethod
    def style_sheet(sheet, widths, row_height, table_name):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B8"
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet.iter_rows(min_row=7):
            sheet.row_dimensions[row[0].row].height = 38 if row[0].row == 7 else row_height
            for cell in row:
                cell.font = Font(name="Arial", size=11, color="17364B")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 7:
                    cell.fill = PatternFill("solid", fgColor="173F56")
                    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        table = Table(displayName=table_name, ref=f"A7:{get_column_letter(sheet.max_column)}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)

    def write_xlsx(self, output_path):
        """Write caller-selected .xlsx path and return it; an existing file is replaced."""
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("output_path must end with .xlsx")
        compact = lambda value: json.dumps(value, ensure_ascii=False)
        workbook = Workbook()
        reference = workbook.active
        reference.title = "Species reference"
        definitions = workbook.create_sheet("Graph definitions")
        for sheet in (reference, definitions):
            for row, text in enumerate([
                "Species structure reference",
                f"{len(self.records)} species converted from supplied SMILES/InChI. RDKit {rdBase.rdkitVersion}.",
                "Explicit H; labels = element + zero-based index + formal charge; parallel lines = bond order.",
                "Base atoms precede appended H. Attached-H counts follow base order; H and H2 use [].",
                "Radical counts are RDKit assignments; no excited state or spin multiplicity is inferred.",
            ], 2):
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
                sheet.cell(row, 1, text).font = Font(name="Arial", size=16 if row == 2 else 11,
                                                   bold=row == 2, color="173F56")
                sheet.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
                sheet.row_dimensions[row].height = 28
        self.write_row(reference, 7, ["Species key", "Formula", "Structure (all atoms)",
            "Base atom indices", "Base bonds (u, v, order)", "Attached H per base atom",
            "Chemical name", "SMILES", "Source SMILES", "Source InChI", "State / representation notes"])
        self.write_row(definitions, 7, ["Species key", "Formula", "Base atoms JSON", "Base bonds JSON",
            "H counts JSON", "All node symbols JSON", "All bonds JSON", "Base formal charges JSON",
            "Chemical name", "State / representation notes", "SMILES", "Source SMILES", "Source InChI",
            "All formal charges JSON", "RDKit radical electrons JSON", "Normalized InChI",
            "Source state annotation"])
        for row, record in enumerate(self.records, 8):
            base_count = len(record["base_atoms"])
            self.write_row(reference, row, [record["species_key"], record["formula"], "",
                ", ".join(f"{i}: {s}" for i, s in enumerate(record["base_atoms"])),
                compact(record["base_bonds"]), compact(record["h_counts"]), record["chemical_name"],
                record["smiles"], record["source_smiles"], record["source_inchi"], record["notes"]])
            self.write_row(definitions, row, [record["species_key"], record["formula"],
                compact(record["base_atoms"]), compact(record["base_bonds"]), compact(record["h_counts"]),
                compact(record["symbols"]), compact(record["bonds"]), compact(record["charges"][:base_count]),
                record["chemical_name"], record["notes"], record["smiles"], record["source_smiles"],
                record["source_inchi"], compact(record["charges"]), compact(record["radicals"]),
                record["normalized_inchi"], record["state_annotation"]])
            picture = Image(io.BytesIO(self.structure_png(record)))
            # Display 420 x 200 px inside a ~453 x 240 px cell, preserving aspect.
            picture.anchor = OneCellAnchor(
                _from=AnchorMarker(col=2, row=row-1, colOff=pixels_to_EMU(12), rowOff=pixels_to_EMU(20)),
                ext=XDRPositiveSize2D(pixels_to_EMU(420), pixels_to_EMU(200)))
            reference.add_image(picture)
        self.style_sheet(reference, [17, 16, 64, 25, 32, 25, 30, 32, 32, 65, 55], 180, "SpeciesReference")
        self.style_sheet(definitions, [17, 16, 28, 36, 24, 44, 58, 28, 30, 55, 32, 32, 65, 30, 34, 65, 26],
                         100, "GraphDefinitions")
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            workbook.save(path)
        finally:
            workbook.close()
        print(f"Saved {len(self.records)} species to {path}")
        return path
